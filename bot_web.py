import os
import time

import pandas as pd
import streamlit as st
from google import genai
from google.genai import types
from openpyxl import Workbook, load_workbook

# El Excel se guarda en la misma carpeta que este archivo
ARCHIVO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "datos.xlsx")
MODELO = "gemini-flash-lite-latest"  # si da problemas, prueba "gemini-flash-latest"


# ---------- Herramientas que la IA puede usar ----------

def abrir():
    if os.path.exists(ARCHIVO):
        return load_workbook(ARCHIVO)
    wb = Workbook()
    wb.save(ARCHIVO)
    return wb


def convertir(v: str):
    """Convierte '25' en número 25 para que Excel lo trate como número."""
    try:
        return float(v) if "." in v else int(v)
    except (ValueError, TypeError):
        return v


def leer_hoja() -> str:
    """Devuelve todo el contenido actual del Excel, fila por fila."""
    filas = [list(f) for f in abrir().active.iter_rows(values_only=True)]
    return str(filas) if filas else "La hoja está vacía."


def escribir_celda(celda: str, valor: str) -> str:
    """Escribe un valor en una celda concreta, por ejemplo B3."""
    wb = abrir()
    wb.active[celda] = convertir(valor)
    wb.save(ARCHIVO)
    return f"{celda} = {valor}"


def anadir_fila(valores: list[str]) -> str:
    """Añade una fila nueva al final de la hoja."""
    wb = abrir()
    wb.active.append([convertir(v) for v in valores])
    wb.save(ARCHIVO)
    return "Fila añadida."


def borrar_fila(numero_fila: int) -> str:
    """Borra la fila indicada (la fila 1 es la primera)."""
    wb = abrir()
    wb.active.delete_rows(numero_fila)
    wb.save(ARCHIVO)
    return f"Fila {numero_fila} borrada."


def enviar(texto: str) -> str:
    """Manda el mensaje a la IA, con reintentos si Google está saturado."""
    for intento in range(1, 4):
        try:
            return st.session_state.chat.send_message(texto).text or "Hecho."
        except PermissionError:
            return "No puedo guardar: cierra datos.xlsx en Excel y vuelve a pedírmelo."
        except Exception as e:
            saturado = "503" in str(e) or "UNAVAILABLE" in str(e)
            if saturado and intento < 3:
                time.sleep(10)
            else:
                return f"Ha habido un error: {e}"


# ---------- Página web ----------

st.set_page_config(page_title="Bot de Excel", page_icon="📊", layout="wide")
st.title("📊 Bot de Excel")

if "chat" not in st.session_state:
    st.session_state.client = genai.Client()  # usa la variable GEMINI_API_KEY
    st.session_state.chat = st.session_state.client.chats.create(
        model=MODELO,
        config=types.GenerateContentConfig(
            tools=[leer_hoja, escribir_celda, anadir_fila, borrar_fila],
            system_instruction=(
                "Eres un asistente que gestiona un archivo Excel. "
                "Lee siempre la hoja antes de modificarla. "
                "Responde en español y de forma breve, confirmando lo que has cambiado."
            ),
        ),
    )
    st.session_state.mensajes = []

# Barra lateral: vista del Excel y botón de descarga
with st.sidebar:
    st.header("Tu Excel")
    filas = [list(f) for f in abrir().active.iter_rows(values_only=True)]
    if filas and any(any(c is not None for c in f) for f in filas):
        df = pd.DataFrame(filas).fillna("").astype(str)
        df.columns = [chr(65 + i) if i < 26 else f"Col{i + 1}" for i in range(len(df.columns))]
        df.index = range(1, len(df) + 1)
        st.dataframe(df, use_container_width=True)
    else:
        st.info("La hoja está vacía. Pídele al bot que cree una tabla.")
    with open(ARCHIVO, "rb") as f:
        st.download_button("⬇️ Descargar datos.xlsx", f, file_name="datos.xlsx")

# Historial del chat
for m in st.session_state.mensajes:
    with st.chat_message(m["rol"]):
        st.markdown(m["texto"])

# Caja para escribir
if texto := st.chat_input("Ej.: crea una tabla con columnas Nombre, Edad y Ciudad"):
    st.session_state.mensajes.append({"rol": "user", "texto": texto})
    with st.chat_message("user"):
        st.markdown(texto)
    with st.chat_message("assistant"):
        with st.spinner("Pensando..."):
            respuesta = enviar(texto)
        st.markdown(respuesta)
    st.session_state.mensajes.append({"rol": "assistant", "texto": respuesta})
    st.rerun()  # recarga para que la tabla de la izquierda muestre los cambios
